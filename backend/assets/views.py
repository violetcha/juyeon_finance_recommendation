from datetime import date, datetime
from functools import lru_cache
from pathlib import Path

from django.conf import settings
from rest_framework import status
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import AllowAny
from rest_framework.response import Response

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


ASSET_CONFIG = {
    'gold': {
        'label': '금',
        'english_label': 'Gold',
        'filename': 'Gold_prices.xlsx',
        'sheet_name': 'Gold',
    },
    'silver': {
        'label': '은',
        'english_label': 'Silver',
        'filename': 'Silver_prices.xlsx',
        'sheet_name': 'Silver',
    },
}


def parse_excel_date(value):
    if value is None:
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    if isinstance(value, (int, float)):
        return from_excel(value).date()

    value = str(value).strip()

    for fmt in ['%Y-%m-%d', '%Y/%m/%d', '%m/%d/%Y', '%m-%d-%Y']:
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            pass

    return None


def parse_number(value):
    if value is None:
        return None

    if isinstance(value, (int, float)):
        return float(value)

    value = str(value).replace(',', '').replace('$', '').strip()

    if value == '':
        return None

    try:
        return float(value)
    except ValueError:
        return None


def parse_query_date(value):
    if not value:
        return None

    try:
        return datetime.strptime(value, '%Y-%m-%d').date()
    except ValueError:
        return None


@lru_cache(maxsize=4)
def load_asset_rows(asset):
    config = ASSET_CONFIG[asset]

    file_path = (
        Path(settings.BASE_DIR)
        / 'assets'
        / 'data'
        / config['filename']
    )

    if not file_path.exists():
        raise FileNotFoundError(f'{config["filename"]} 파일을 찾을 수 없습니다: {file_path}')

    workbook = load_workbook(file_path, data_only=True)
    worksheet = workbook[config['sheet_name']]

    rows = []

    for row in worksheet.iter_rows(min_row=2, values_only=True):
        raw_date, close, volume, open_price, high, low = row[:6]

        parsed_date = parse_excel_date(raw_date)

        if parsed_date is None:
            continue

        close_value = parse_number(close)

        if close_value is None:
            continue

        rows.append({
            'date': parsed_date.isoformat(),
            'close': close_value,
            'volume': parse_number(volume),
            'open': parse_number(open_price),
            'high': parse_number(high),
            'low': parse_number(low),
        })

    rows.sort(key=lambda item: item['date'])

    return rows


def build_summary(rows):
    if not rows:
        return None

    first = rows[0]
    latest = rows[-1]

    first_close = first['close']
    latest_close = latest['close']

    change = latest_close - first_close
    change_rate = 0

    if first_close:
        change_rate = (change / first_close) * 100

    close_values = [row['close'] for row in rows]

    return {
        'first_date': first['date'],
        'latest_date': latest['date'],
        'first_close': round(first_close, 3),
        'latest_close': round(latest_close, 3),
        'change': round(change, 3),
        'change_rate': round(change_rate, 2),
        'min_close': round(min(close_values), 3),
        'max_close': round(max(close_values), 3),
    }


@api_view(['GET'])
@permission_classes([AllowAny])
def asset_prices(request):
    asset = request.GET.get('asset', 'gold').lower()
    start_date_param = request.GET.get('start_date')
    end_date_param = request.GET.get('end_date')

    if asset not in ASSET_CONFIG:
        return Response(
            {'message': 'asset은 gold 또는 silver만 가능합니다.'},
            status=status.HTTP_400_BAD_REQUEST
        )

    start_date = parse_query_date(start_date_param)
    end_date = parse_query_date(end_date_param)

    if start_date_param and start_date is None:
        return Response(
            {'message': '시작일 형식이 올바르지 않습니다. YYYY-MM-DD 형식으로 입력해주세요.'},
            status=status.HTTP_400_BAD_REQUEST
        )

    if end_date_param and end_date is None:
        return Response(
            {'message': '종료일 형식이 올바르지 않습니다. YYYY-MM-DD 형식으로 입력해주세요.'},
            status=status.HTTP_400_BAD_REQUEST
        )

    if start_date and end_date and start_date > end_date:
        return Response(
            {'message': '시작일은 종료일보다 늦을 수 없습니다.'},
            status=status.HTTP_400_BAD_REQUEST
        )

    try:
        rows = load_asset_rows(asset)
    except FileNotFoundError as error:
        return Response(
            {'message': str(error)},
            status=status.HTTP_500_INTERNAL_SERVER_ERROR
        )

    filtered_rows = []

    for row in rows:
        row_date = datetime.strptime(row['date'], '%Y-%m-%d').date()

        if start_date and row_date < start_date:
            continue

        if end_date and row_date > end_date:
            continue

        filtered_rows.append(row)

    if not filtered_rows:
        return Response({
            'asset': asset,
            'asset_label': ASSET_CONFIG[asset]['label'],
            'count': 0,
            'message': '선택한 기간에 해당하는 데이터가 없습니다.',
            'rows': [],
            'summary': None,
        })

    return Response({
        'asset': asset,
        'asset_label': ASSET_CONFIG[asset]['label'],
        'english_label': ASSET_CONFIG[asset]['english_label'],
        'count': len(filtered_rows),
        'rows': filtered_rows,
        'summary': build_summary(filtered_rows),
    })